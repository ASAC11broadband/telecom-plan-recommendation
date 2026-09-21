"""코사인 유사도 기반 콘텐츠 추천을 원본 합성 가입이력 100명으로 평가한다.

    python run_cosine_test.py
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import openpyxl

from agent.cosine_recommendation import prepare_plan_catalog, recommend_by_cosine
from agent.segmentation import ROOT, evaluate, holdout_100, load_interactions


OUTPUT_DIR = ROOT / "data" / "synthetic_original"
DEFAULT_OUTPUT = OUTPUT_DIR / "cosine_recommendation_evaluation.xlsx"
DEFAULT_JSONL = OUTPUT_DIR / "cosine_recommendation_evaluation.jsonl"


def _result_rows(test_rows, catalog, limit: int) -> list[dict[str, Any]]:
    results = []
    for _, row in test_rows.iterrows():
        recommendation = recommend_by_cosine(catalog, row, limit=limit)
        plan_ids = [item["plan_id"] for item in recommendation["ranked"]]
        actual_id = str(row["source_plan_id"])
        try:
            rank = plan_ids.index(actual_id) + 1
        except ValueError:
            rank = None
        results.append(
            {
                "customer_id": str(row["customer_id"]),
                "data_gb_month": float(row["data_gb_month"]),
                "budget_krw": int(row["budget_krw"]),
                "candidate_count": recommendation["candidate_count"],
                "actual_plan_id": actual_id,
                "actual_plan_name": str(row["plan_name"]),
                "ranked": recommendation["ranked"],
                "hit_at_5": rank is not None,
                "rank": rank,
            }
        )
    return results


def write_outputs(results: list[dict[str, Any]], summary: dict[str, Any]) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    DEFAULT_JSONL.write_text(
        "\n".join(json.dumps(row, ensure_ascii=False) for row in results) + "\n",
        encoding="utf-8",
    )
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "코사인 유사도 추천"
    headers = [
        "customer_id", "data_gb_month", "budget_krw", "candidate_count",
        "actual_plan_id", "actual_plan_name", "top1_plan_name", "top1_cosine_similarity",
        "top5_plan_ids", "top5_plan_names", "hit_at_5", "rank",
    ]
    sheet.append(headers)
    for result in results:
        ranked = result["ranked"]
        sheet.append(
            [
                result["customer_id"], result["data_gb_month"], result["budget_krw"], result["candidate_count"],
                result["actual_plan_id"], result["actual_plan_name"],
                ranked[0]["plan_name"] if ranked else "",
                ranked[0]["cosine_similarity"] if ranked else "",
                " | ".join(item["plan_id"] for item in ranked),
                " | ".join(item["plan_name"] for item in ranked),
                "Y" if result["hit_at_5"] else "N", result["rank"] or "",
            ]
        )
    for column, width in {1: 14, 6: 36, 7: 36, 9: 44, 10: 100}.items():
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"

    overview = workbook.create_sheet("평가 요약")
    overview.append(["metric", "value"])
    for key, value in summary.items():
        overview.append([key, value])
    overview.column_dimensions["A"].width = 46
    overview.column_dimensions["B"].width = 28
    workbook.save(DEFAULT_OUTPUT)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--limit", type=int, default=5)
    args = parser.parse_args()

    interactions = load_interactions()
    _, test = holdout_100(interactions, seed=args.seed, size=100)
    results = _result_rows(test, prepare_plan_catalog(), args.limit)
    metrics = evaluate(results)
    raw_customer_rows = int(interactions.attrs["raw_customer_rows"])
    summary: dict[str, Any] = {
        "raw_customer_rows": raw_customer_rows,
        "joined_interaction_rows": len(interactions),
        "plan_join_coverage": round(len(interactions) / raw_customer_rows, 4),
        "holdout_users": len(results),
        "catalog_candidates": 2235,
        "ranking_method": "hard requirement filter -> weighted cosine similarity (70%) -> resource/budget/OTT correction (30%)",
        "vector_features": "data, voice, sms, unlimited needs, price sensitivity, OTT need",
        "mean_candidates_after_hard_filter": round(sum(row["candidate_count"] for row in results) / len(results), 1),
        "top1_accuracy": round(sum(row["rank"] == 1 for row in results) / len(results), 4),
        **metrics,
    }
    write_outputs(results, summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"결과: {DEFAULT_OUTPUT}")


if __name__ == "__main__":
    main()
