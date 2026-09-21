"""원본 합성 가입 이력으로 100명 홀드아웃 추천 실험을 재현한다.

    python -m experiments.run_segment_test

초기안은 ``세그먼트 소속 확률 × 세그먼트별 인기 점수``이고, 전환안은
``프로필 유사 이웃 -> 거리 가중 요금제 Top-5``이다. 두 방식 모두 학습 데이터에만
접근하고, 분리한 100명의 ``source_plan_id``를 정답으로 평가한다.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Callable

import openpyxl

from agent.segmentation import (
    ROOT,
    evaluate,
    fit_model,
    holdout_100,
    load_interactions,
    neighbor_recommend,
    segment_popularity_recommend,
)


OUTPUT_DIR = ROOT / "data" / "synthetic_original"
DEFAULT_OUTPUT = OUTPUT_DIR / "segment_recommendation_comparison.xlsx"
DEFAULT_JSONL = OUTPUT_DIR / "segment_recommendation_comparison.jsonl"


def _result_rows(
    test_rows,
    recommend: Callable[..., dict[str, Any]],
    model,
    limit: int,
) -> list[dict[str, Any]]:
    results = []
    for _, row in test_rows.iterrows():
        recommendation = recommend(model, row.to_frame().T, limit=limit)
        plan_ids = [item["plan_id"] for item in recommendation["ranked"]]
        actual_id = str(row["source_plan_id"])
        try:
            rank = plan_ids.index(actual_id) + 1
        except ValueError:
            rank = None
        results.append(
            {
                "customer_id": str(row["customer_id"]),
                "segment": recommendation["segment"],
                "segment_probabilities": recommendation.get("segment_probabilities", []),
                "age": int(row["age"]),
                "data_gb_month": float(row["data_gb_month"]),
                "budget_krw": int(row["budget_krw"]),
                "actual_plan_id": actual_id,
                "actual_plan_name": str(row["plan_name"]),
                "ranked": recommendation["ranked"],
                "hit_at_5": rank is not None,
                "rank": rank,
            }
        )
    return results


def _sheet_row(result: dict[str, Any]) -> list[Any]:
    ranked = result["ranked"]
    probabilities = result["segment_probabilities"]
    segment_mix = " | ".join(f"S{index + 1} {probability:.1%}" for index, probability in enumerate(probabilities))
    return [
        result["customer_id"],
        f"S{result['segment'] + 1}",
        max(probabilities) if probabilities else "",
        segment_mix,
        result["age"],
        result["data_gb_month"],
        result["budget_krw"],
        result["actual_plan_id"],
        result["actual_plan_name"],
        ranked[0]["plan_name"] if ranked else "",
        " | ".join(item["plan_id"] for item in ranked),
        " | ".join(item["plan_name"] for item in ranked),
        "Y" if result["hit_at_5"] else "N",
        result["rank"] or "",
    ]


def _add_result_sheet(sheet, title: str, results: list[dict[str, Any]]) -> None:
    sheet.title = title
    headers = [
        "customer_id", "dominant_segment", "dominant_segment_probability", "segment_mix",
        "age", "data_gb_month", "budget_krw",
        "actual_plan_id", "actual_plan_name", "top1_plan_name", "top5_plan_ids",
        "top5_plan_names", "hit_at_5", "rank",
    ]
    sheet.append(headers)
    for result in results:
        sheet.append(_sheet_row(result))
    for column, width in {1: 14, 4: 72, 9: 36, 10: 36, 11: 44, 12: 100}.items():
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"


def write_outputs(
    baseline: list[dict[str, Any]],
    personalized: list[dict[str, Any]],
    summary: dict[str, Any],
    output_path: Path,
    jsonl_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    paired = [
        {"customer_id": old["customer_id"], "soft_segment_popularity_baseline": old, "profile_neighbor_personalized": new}
        for old, new in zip(baseline, personalized)
    ]
    jsonl_path.write_text(
        "\n".join(json.dumps(row, ensure_ascii=False) for row in paired) + "\n",
        encoding="utf-8",
    )

    workbook = openpyxl.Workbook()
    _add_result_sheet(workbook.active, "소프트 세그먼트 초기안", baseline)
    _add_result_sheet(workbook.create_sheet(), "개인 이웃 전환안", personalized)
    overview = workbook.create_sheet("비교 요약")
    overview.append(["metric", "value"])
    for key, value in summary.items():
        overview.append([key, value])
    overview.column_dimensions["A"].width = 42
    overview.column_dimensions["B"].width = 20
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--clusters", type=int, default=6)
    parser.add_argument("--neighbors", type=int, default=100)
    parser.add_argument("--limit", type=int, default=5)
    # 1차 합성은 2026-08-12 스냅샷, 2차는 현행 카탈로그로 만들었다. 둘 다 같은 코드로 재평가한다.
    parser.add_argument("--customers", type=Path, default=None)
    parser.add_argument("--plans", type=Path, default=None)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    paths = {k: v for k, v in (("customers_path", args.customers), ("plans_path", args.plans)) if v}
    interactions = load_interactions(**paths)
    train, test = holdout_100(interactions, seed=args.seed, size=100)
    model = fit_model(train, clusters=args.clusters, neighbors=args.neighbors)
    baseline = _result_rows(test, segment_popularity_recommend, model, args.limit)
    personalized = _result_rows(test, neighbor_recommend, model, args.limit)

    baseline_metrics = evaluate(baseline)
    personalized_metrics = evaluate(personalized)
    raw_customer_rows = int(interactions.attrs["raw_customer_rows"])
    summary: dict[str, Any] = {
        "raw_customer_rows": raw_customer_rows,
        "joined_interaction_rows": len(interactions),
        "plan_join_coverage": round(len(interactions) / raw_customer_rows, 4),
        "holdout_users": len(test),
        "cluster_count": args.clusters,
        "neighbor_count": min(args.neighbors, len(train)),
        "segment_model": "KMeans 6-centroid distance-softened membership",
        "baseline_method": "segment membership probability × smoothed segment plan popularity Top-5",
        "personalized_method": "profile-neighbor distance-weighted plan Top-5",
    }
    summary.update({f"baseline_{key}": value for key, value in baseline_metrics.items()})
    summary.update({f"personalized_{key}": value for key, value in personalized_metrics.items()})
    summary["hit_at_5_lift"] = round(personalized_metrics["hit_at_5"] - baseline_metrics["hit_at_5"], 4)
    summary["mrr_at_5_lift"] = round(personalized_metrics["mrr_at_5"] - baseline_metrics["mrr_at_5"], 4)
    write_outputs(baseline, personalized, summary, args.out, args.out.with_suffix(".jsonl"))
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"결과: {args.out}")


if __name__ == "__main__":
    main()
