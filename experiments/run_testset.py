# -*- coding: utf-8 -*-
"""정답지 문항을 파이프라인에 태워 Calc_precision_recall.py 가 읽는 wide 포맷으로 저장한다.

    python -m experiments.run_testset                       # 100문항 전체 -> data/recommend_results.xlsx
    python -m experiments.run_testset --limit 5             # 스모크 테스트
    python -m experiments.run_testset --out results_4o.xlsx --workers 8
    python -m experiments.run_testset --ids L2-28,L3-09 --workers 1   # 429 등으로 실패한 문항만 재실행

--ids 로 돌리면 기존 --out 파일을 읽어 그 문항 행만 덮어쓰고 나머지는 그대로 둔다.
"""

from __future__ import annotations

import argparse
import os
from concurrent.futures import ThreadPoolExecutor

import openpyxl
from langchain_core.messages import HumanMessage
from openpyxl import Workbook

from agent.graph import graph

ANSWER_KEY_PATH = "data/test_cases_정답지 - 복사본.xlsx"
SHEET_NAME = "테스트케이스_v2"
MAX_RANK = 5


def load_questions(path: str) -> list[tuple]:
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET_NAME]
    return [
        (ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value)
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=1).value is not None
    ]


def answer_one(question: str) -> list[dict]:
    """추천 결과를 정답지와 같은 항목 형태(name/monthly/discounted/carrier_type)로."""
    state = graph.invoke({"messages": [HumanMessage(content=question)]})
    by_id = {c["plan_id"]: c for c in state.get("candidates", [])}
    items = []
    for plan in state.get("ranked", [])[:MAX_RANK]:
        c = by_id.get(plan.plan_id)
        if c is None:
            continue
        items.append({
            "name": c["plan_name"],
            "monthly_fee": c["monthly_fee"],
            "discounted_fee": c["discounted_fee"],
            "carrier_type": c["carrier_type"],
        })
    return items


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--answer-key", default=ANSWER_KEY_PATH)
    ap.add_argument("--out", default="data/recommend_results.xlsx")
    ap.add_argument("--limit", type=int)
    ap.add_argument("--ids", help="쉼표로 구분한 문항 id만 재실행. 나머지 문항은 기존 --out 파일 값을 유지한다")
    ap.add_argument("--workers", type=int, default=4)
    args = ap.parse_args()

    all_questions = load_questions(args.answer_key)
    if args.ids:
        wanted = {x.strip() for x in args.ids.split(",") if x.strip()}
        questions = [q for q in all_questions if q[0] in wanted]
        missing = wanted - {q[0] for q in questions}
        if missing:
            ap.error(f"정답지에 없는 id: {sorted(missing)}")
    else:
        questions = all_questions[: args.limit]

    def work(row):
        qid, level, question = row
        try:
            items = answer_one(question)
        except Exception as exc:  # 한 문항 실패로 배치 전체를 버리지 않는다
            print(f"[ERROR] {qid}: {exc}")
            items = []
        print(f"{qid} ({len(items)}개) {question}")
        return qid, level, question, items

    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        results = list(pool.map(work, questions))

    # 기존 결과 파일이 있으면 그 위에 이번 실행분만 덮어쓴다
    rows: dict = {}
    if os.path.exists(args.out):
        ws_old = openpyxl.load_workbook(args.out, data_only=True)[SHEET_NAME]
        for r in range(2, ws_old.max_row + 1):
            values = [ws_old.cell(row=r, column=c).value for c in range(1, 4 + MAX_RANK * 4)]
            if values[0] is not None:
                rows[values[0]] = values

    for qid, level, question, items in results:
        row = [qid, level, question]
        for i in range(MAX_RANK):
            item = items[i] if i < len(items) else None
            row += [item["name"], item["monthly_fee"], item["discounted_fee"], item["carrier_type"]] if item else [None] * 4
        rows[qid] = row

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    header = ["id", "level", "question"]
    for i in range(1, MAX_RANK + 1):
        header += [f"ans{i}_name", f"ans{i}_monthly_fee", f"ans{i}_discounted_fee", f"ans{i}_carrier_type"]
    ws.append(header)
    for qid, _, _ in all_questions:  # 정답지 순서 유지
        if qid in rows:
            ws.append(rows[qid])
    wb.save(args.out)
    print(f"\n완료: {args.out} (이번 실행 {len(results)}문항, 파일 전체 {len(rows)}문항)")


if __name__ == "__main__":
    main()
