"""콘텐츠 기반 추천 두 판을 합성 가입이력 100명으로 비교한다.

    python -m experiments.run_cosine_test

`v1 코사인`은 교과서대로고, `v2 비대칭 충족도`는 그 결함(가격을 벡터에 넣음 / 질의에
없는 축이 벌점 / 과잉 제공에 벌점)을 고친 판이다. 근거는 agent.cosine_recommendation
모듈 설명에 있다.

Hit@5 는 **절대 성능이 아니다.** 합성 가입이력은 요금제에서 사용자를 역산해 만들어져
생성 규칙을 3줄로 뒤집으면 Hit@5 0.905 가 나온다. 그래서 이 숫자는 두 판의 상대 비교와
"생성 규칙을 얼마나 복원했나"로만 읽어야 한다. 품질 판단은 답지가 필요 없는 지표
(파레토 비지배율·가격 백분위·예산 단조성)로 따로 한다.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Callable

import openpyxl

from agent.cosine_recommendation import (prepare_plan_catalog, recommend_by_cosine,
                                         recommend_by_coverage)
from agent.segmentation import ROOT, evaluate, holdout_100, load_interactions

OUTPUT_DIR = ROOT / "data" / "synthetic_original"
DEFAULT_OUTPUT = OUTPUT_DIR / "cosine_recommendation_evaluation.xlsx"

ARMS: dict[str, tuple[str, Callable, str]] = {
    "cosine": ("v1 코사인 유사도", recommend_by_cosine, "cosine_similarity"),
    "coverage": ("v2 비대칭 충족도", recommend_by_coverage, "coverage"),
}


def _result_rows(test_rows, catalog, limit: int, recommend: Callable, score_key: str) -> list[dict[str, Any]]:
    results = []
    for _, row in test_rows.iterrows():
        try:
            recommendation = recommend(catalog, row, limit=limit)
        except ValueError:  # 조건을 동시에 만족하는 후보가 없는 사용자
            recommendation = {"candidate_count": 0, "ranked": []}
        plan_ids = [item["plan_id"] for item in recommendation["ranked"]]
        actual_id = str(row["source_plan_id"])
        rank = plan_ids.index(actual_id) + 1 if actual_id in plan_ids else None
        results.append({
            "customer_id": str(row["customer_id"]),
            "data_gb_month": float(row["data_gb_month"]) if row["data_gb_month"] == row["data_gb_month"] else None,
            "budget_krw": int(row["budget_krw"]),
            "candidate_count": recommendation["candidate_count"],
            "actual_plan_id": actual_id,
            "actual_plan_name": str(row["plan_name"]),
            "ranked": recommendation["ranked"],
            "top_score": recommendation["ranked"][0][score_key] if recommendation["ranked"] else None,
            "hit_at_5": rank is not None,
            "rank": rank,
        })
    return results


def _sheet(workbook, title: str, results: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(["customer_id", "data_gb_month", "budget_krw", "candidate_count",
                  "actual_plan_id", "actual_plan_name", "top1_plan_name", "top1_score",
                  "top5_plan_ids", "top5_plan_names", "hit_at_5", "rank"])
    for result in results:
        ranked = result["ranked"]
        sheet.append([
            result["customer_id"], result["data_gb_month"], result["budget_krw"],
            result["candidate_count"], result["actual_plan_id"], result["actual_plan_name"],
            ranked[0]["plan_name"] if ranked else "", result["top_score"] or "",
            " | ".join(item["plan_id"] for item in ranked),
            " | ".join(item["plan_name"] for item in ranked),
            "Y" if result["hit_at_5"] else "N", result["rank"] or "",
        ])
    for column, width in {1: 14, 6: 36, 7: 36, 9: 44, 10: 100}.items():
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"


def write_outputs(per_arm: dict[str, list[dict[str, Any]]], summary: dict[str, Any], out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    for key, results in per_arm.items():
        out.with_name(f"{out.stem}_{key}.jsonl").write_text(
            "\n".join(json.dumps(row, ensure_ascii=False) for row in results) + "\n", encoding="utf-8")

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for key, results in per_arm.items():
        _sheet(workbook, ARMS[key][0], results)
    overview = workbook.create_sheet("비교 요약")
    overview.append(["metric", "value"])
    for key, value in summary.items():
        overview.append([key, value])
    overview.column_dimensions["A"].width = 46
    overview.column_dimensions["B"].width = 30
    workbook.save(out)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--limit", type=int, default=5)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    interactions = load_interactions()
    _, test = holdout_100(interactions, seed=args.seed, size=100)
    catalog = prepare_plan_catalog()

    per_arm, summary = {}, {}
    raw_rows = int(interactions.attrs["raw_customer_rows"])
    summary.update({
        "raw_customer_rows": raw_rows,
        "joined_interaction_rows": len(interactions),
        "plan_join_coverage": round(len(interactions) / raw_rows, 4),
        "holdout_users": len(test),
        "catalog_candidates": len(catalog),
    })
    for key, (label, recommend, score_key) in ARMS.items():
        results = _result_rows(test, catalog, args.limit, recommend, score_key)
        per_arm[key] = results
        n = len(results)
        summary[f"{key}_label"] = label
        summary[f"{key}_mean_candidates"] = round(sum(r["candidate_count"] for r in results) / n, 1)
        summary[f"{key}_top1_accuracy"] = round(sum(r["rank"] == 1 for r in results) / n, 4)
        summary[f"{key}_unique_top5_count"] = len({" ".join(i["plan_id"] for i in r["ranked"]) for r in results})
        summary.update({f"{key}_{m}": v for m, v in evaluate(results).items()})

    write_outputs(per_arm, summary, args.out)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"결과: {args.out}")


if __name__ == "__main__":
    main()
