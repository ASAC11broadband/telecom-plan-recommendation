# -*- coding: utf-8 -*-
"""가중치 명세(arm)별 추천 품질 비교 — 후보군을 얼려놓고 순위만 다시 매긴다.

**지금은 실행되지 않는다(보관용).** de85a04 에서 `agent.agents.recommend._dedupe_by_name` 이
`_dedupe_identical_offers` 로 바뀌고 `agent.mcda._finite_voice` 가 없어졌다. 중복 정리 규칙 자체가
달라져서 이름만 고쳐서는 data/weight_arms_* 의 수치(난수 0.292 / Ridge 0.370)가 재현되지 않는다.

가중치는 `evaluate_mcda` 에만 들어간다. 후보군 생성(필터·비교목표)까지는 가중치와
무관하고, 그 앞의 프로파일링은 LLM 이라 돌릴 때마다 달라진다. arm 마다 파이프라인을
통째로 다시 태우면 그 비결정성이 가중치 효과와 섞인다(빈 응답 22 vs 24 가 P/R 을
±0.02 흔든다). 그래서 프로파일링은 한 번만 태워 캐시하고, arm 은 순위만 다시 매긴다.

    python -m experiments.run_weight_arms                 # 캐시 있으면 LLM 안 태움
    python -m experiments.run_weight_arms --refresh       # 프로파일·가중치 캐시 재생성

산출: data/weight_arms_results.csv (arm x 문항 P/R)
"""

from __future__ import annotations

import argparse
import json
import sys
import math
import re
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from langchain_core.messages import HumanMessage

from experiments.Calc_precision_recall import count_intersection, load_wide_rows
from experiments import mcda_original  # eb0e8a2^ 시점 agent/mcda.py — 감마 난수 가중치 SMAA-2
from agent.agents.profiling import profiling_node
from agent.agents.recommend import _apply_comparison, _dedupe_by_name, _resolve_reference
from agent.data import filter_candidates
from agent.mcda import _finite_data, _finite_voice, _minmax
from agent.schemas import UserProfile

DATA = Path("data")
ANSWER_KEY = DATA / "test_cases_정답지.xlsx"
PROFILE_CACHE = DATA / "testset_profiles.json"
ARM_CACHE = DATA / "weight_arms.json"
SHAP_WEIGHTS = DATA / "weight_bootstrap_shap.json"
OUT_CSV = DATA / "weight_arms_results.csv"
OUT_DETAIL = DATA / "weight_arms_detail.xlsx"
OUT_BENCH = DATA / "benchmark"
SHEET = "테스트케이스_v2"
TOP_K = 5

# ---------------------------------------------------------------- 가중치 arm

RAW = {"cost_month": "raw", "data_total_gb": "raw"}
TRANSFORMED = {"cost_month": "sqrt", "data_total_gb": "log1p"}
AXES7 = ["cost_month", "data_total_gb", "qos", "tether_gb", "voice_min", "sms_cnt", "ott_value"]
AXES6 = [a for a in AXES7 if a != "sms_cnt"]
PROD = {"cost_month": "price", "data_total_gb": "data", "qos": "qos",
        "tether_gb": "tethering", "voice_min": "voice", "ott_value": "benefit"}
# arm 별: (설명, 축, 변환, 문자 사후제거 여부, 부트스트랩 n)
ARM_SPECS = {
    "B_ridge7_raw": ("Ridge · 7축 raw · 문자 사후제거", AXES7, RAW, True, 300),
    "C_ridge6_raw": ("Ridge · 6축 raw · 문자 사전제거", AXES6, RAW, False, 300),
    "D_ridge6_tf": ("Ridge · 6축 sqrt/log1p · 문자 사전제거", AXES6, TRANSFORMED, False, 300),
    "D60_ridge6_tf": ("D 와 같되 부트스트랩 60회(이전 기본값)", AXES6, TRANSFORMED, False, 60),
}
TRANSFORMS = {"raw": lambda v: v, "log1p": np.log1p, "sqrt": np.sqrt}


def _mvno_features() -> pd.DataFrame:
    """logshare_shap_weights / linear_spec_diagnostic 과 같은 모집단·피처."""
    plans = pd.read_csv(DATA / "통신요금제_통합데이터_최종.csv", dtype={"plan_id": str})
    benefits = pd.read_csv(DATA / "통신요금제_혜택상세_최종.csv", dtype={"plan_id": str})
    prices = pd.read_csv(DATA / "ott_prices.csv")
    price_map = dict(zip(prices.service, prices.monthly_won.astype(int)))

    mno = ["SKT", "KT", "LG U+", "LGU+"]
    d = plans[(plans.carrier_type == "MVNO") & (~plans.mvno_brand.isin(mno))].copy()

    def qos_mbps(value):
        if pd.isna(value):
            return 0.0
        m = re.match(r"([\d.]+)\s*(kbps|mbps)", str(value).strip().lower().replace(" ", ""))
        if not m:
            return 0.0
        n = float(m.group(1))
        return n / 1000 if m.group(2) == "kbps" else n

    extra = benefits[benefits.benefit_category == "추가데이터"].copy()
    extra["gb"] = extra.benefit_name.fillna("").str.extract(r"([\d.]+)\s*GB", flags=re.I)[0].astype(float)
    d["extra_gb"] = d.plan_id.map(extra.groupby("plan_id").gb.sum()).fillna(0)
    d["cost_month"] = d.discounted_fee
    d["data_total_gb"] = np.where(
        d.data_unlimited, 300,
        (d.data_gb.fillna(d.daily_data_gb.fillna(0) * 30) + d.extra_gb).clip(0, 300))
    d["qos"] = d.data_throttle_speed.map(qos_mbps)
    d["tether_gb"] = d.tethering_gb.fillna(0)
    d["voice_min"] = np.where(d.voice_unlimited, 1000, d.voice_minutes.fillna(0)).clip(0, 1000)
    d["sms_cnt"] = np.where(d.sms_unlimited, 1000, d.sms_count.fillna(0)).clip(0, 1000)
    ott = benefits[benefits.benefit_service.isin(price_map)]  # 카테고리명이 바뀌어(2026-09-16 크롤러 수정) 서비스명 매칭으로 변경
    services = d.plan_id.map(ott.groupby("plan_id")["benefit_service"].apply(lambda v: sorted(set(v.dropna()))))
    d["ott_value"] = services.apply(lambda s: sum(price_map.get(x, 0) for x in s) if isinstance(s, list) else 0)
    return d


def build_arms() -> dict:
    """Ridge arm 별 가중치 표본. 부트스트랩이라 몇 분 걸린다."""
    from sklearn.linear_model import RidgeCV
    from sklearn.pipeline import make_pipeline
    from sklearn.preprocessing import StandardScaler

    alphas = np.logspace(-3, 3, 25)
    mvno = _mvno_features()
    y = np.log((mvno.subscriber_count / mvno.subscriber_count.sum()).clip(lower=1e-12).values)
    brand = pd.get_dummies(mvno.mvno_brand, prefix="brand", dummy_na=True).astype(float)

    def weights(axes, transform, post_drop, idx=None):
        X = pd.concat([mvno[axes].reset_index(drop=True), brand.reset_index(drop=True)], axis=1)
        for ax, t in transform.items():
            X[ax] = TRANSFORMS[t](X[ax].values)
        yv = y if idx is None else y[idx]
        Xs = X if idx is None else X.iloc[idx]
        fit = make_pipeline(StandardScaler(), RidgeCV(alphas=alphas)).fit(Xs, yv)
        c = pd.Series(fit[-1].coef_, index=X.columns)[axes].abs()
        if post_drop:  # 7축 회귀 후 문자만 빼고 재정규화 — 부풀려진 통화 계수가 그대로 남는다
            c = c.drop("sms_cnt")
        return (c / c.sum()).rename(PROD)

    arms = {}
    for key, (label, axes, transform, post_drop, n) in ARM_SPECS.items():
        rng = np.random.default_rng(0)
        draws = []
        for _ in range(n):
            i = rng.choice(len(mvno), len(mvno), replace=True)
            draws.append(weights(axes, transform, post_drop, i))
        df = pd.DataFrame(draws)
        base = weights(axes, transform, post_drop)
        arms[key] = {
            "label": label,
            "criteria": list(df.columns),
            "base": base.to_dict(),
            "transform": {PROD[k]: v for k, v in transform.items()},
            "vectors": df.values.tolist(),
        }
        print(f"  {key:<16} {n}세트 · {base.round(3).to_dict()}")
    return arms


def load_arms(refresh: bool) -> dict:
    """arm0(가격 단독)·A(현행 SHAP)는 만들 게 없고, Ridge arm 만 부트스트랩한다."""
    if ARM_CACHE.exists() and not refresh:
        arms = json.loads(ARM_CACHE.read_text(encoding="utf-8"))
    else:
        arms = build_arms()
        ARM_CACHE.write_text(json.dumps(arms, ensure_ascii=False), encoding="utf-8")

    shap = json.loads(SHAP_WEIGHTS.read_text(encoding="utf-8"))
    ordered = {
        "R_random_gamma": {
            "label": "SMAA-2 원본 - 감마 난수 가중치 1,200세트 · 10축",
            "criteria": list(mcda_original.CRITERIA), "transform": {}, "base": None,
            "vectors": None,  # 후보군마다 다시 뽑는다(우선순위·비교목표 alpha 가산)
        },
        "arm0_price_only": {
            "label": "가격 단독 - 정답지 생성 규칙과 같은 기준",
            "criteria": ["price"], "transform": {}, "base": {"price": 1.0}, "vectors": [[1.0]],
        },
        "A_shap7": {
            "label": "LGBM+SHAP · 7축 · 현행 main",
            "criteria": shap["criteria"], "transform": {"price": "log1p", "data": "log1p"},
            "base": None, "vectors": shap["vectors"],
        },
    }
    ordered.update(arms)
    return ordered


# ---------------------------------------------------------------- 후보군 동결

def load_questions() -> list[tuple]:
    ws = openpyxl.load_workbook(ANSWER_KEY, data_only=True)[SHEET]
    return [(ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value)
            for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value is not None]


def profile_of(question: str) -> dict:
    out = profiling_node({"messages": [HumanMessage(content=question)]}, {})
    return out["profile"].model_dump(mode="json")


def load_profiles(questions: list[tuple], refresh: bool, workers: int) -> dict:
    if PROFILE_CACHE.exists() and not refresh:
        return json.loads(PROFILE_CACHE.read_text(encoding="utf-8"))

    def work(row):
        qid, _, question = row
        try:
            return qid, profile_of(question)
        except Exception as exc:
            print(f"[ERROR] {qid}: {exc}")
            return qid, None

    with ThreadPoolExecutor(max_workers=workers) as pool:
        profiles = dict(pool.map(work, questions))

    failed = [k for k, v in profiles.items() if v is None]
    if failed:  # 429 등 — 실패분만 직렬 재시도
        print(f"재시도 {len(failed)}문항")
        for qid, _, question in questions:
            if qid in failed:
                profiles[qid] = profile_of(question)
    PROFILE_CACHE.write_text(json.dumps(profiles, ensure_ascii=False), encoding="utf-8")
    return profiles


def candidates_of(profile_dict: dict) -> list[dict]:
    """recommend_node 와 같은 순서로 후보를 좁힌다. 여기까지는 가중치와 무관하다."""
    profile = UserProfile(**profile_dict)
    reference, question = _resolve_reference(profile)
    if question:
        return []
    rows = filter_candidates(profile.model_dump())
    rows = _apply_comparison(rows, reference, profile.comparison_goals or [])
    return _dedupe_by_name(rows)


# ---------------------------------------------------------------- arm 별 순위

_PRIORITY_UNIT = 10.0
_GOAL_UNIT = 12.0
GOAL_CRITERION = {"cheaper": "price", "more_data": "data", "faster_qos": "qos"}


def utility_rows(candidates: list[dict], criteria: list[str], transform: dict) -> list[list[float]]:
    """mcda._utility_rows 와 같되 축 집합·변환을 arm 마다 바꾼다."""
    funcs = {"raw": lambda v: v, "log1p": math.log1p, "sqrt": math.sqrt}
    data_max = max((float(p.get("data_gb") or 0) for p in candidates), default=1.0) or 1.0
    voice_max = max((float(p.get("voice_minutes") or 0) for p in candidates), default=1.0) or 1.0
    sms_max = max((float(p.get("sms_count") or 0) for p in candidates), default=1.0) or 1.0

    def tf(name, value):
        return funcs[transform.get(name, "raw")](value)

    columns = {
        "price": _minmax([tf("price", float(p.get("discounted_fee") or 0)) for p in candidates], cost=True),
        "data": _minmax([tf("data", _finite_data(p, data_max)) for p in candidates]),
        "qos": _minmax([float(p.get("qos_mbps") or 0) for p in candidates]),
        "benefit": _minmax([float(len(p.get("included_benefits") or [])) for p in candidates]),
        "voice": _minmax([_finite_voice(p, voice_max) for p in candidates]),
        "sms": _minmax([(sms_max * 1.25 if p.get("sms_unlimited") else float(p.get("sms_count") or 0))
                        for p in candidates]),
        "tethering": _minmax([float(p.get("tethering_gb") or 0) for p in candidates]),
    }
    return [[columns[name][i] for name in criteria] for i in range(len(candidates))]


def boosted(vector: list[float], criteria: list[str], priorities, goals) -> list[float]:
    out = list(vector)
    unit_priority = _PRIORITY_UNIT / len(criteria)
    unit_goal = _GOAL_UNIT / len(criteria)
    priorities = priorities or []
    for position, name in enumerate(priorities):
        if name in criteria:
            out[criteria.index(name)] += max(0, len(priorities) - position) * unit_priority
    for goal in goals or []:
        name = GOAL_CRITERION.get(goal)
        if name in criteria:
            out[criteria.index(name)] += unit_goal
    total = sum(out)
    return [value / total for value in out]


def smaa_scores(candidates: list[dict], arm: dict, profile: dict) -> list[tuple]:
    """후보 전체를 SMAA-2 기대순위로 정렬해 (요금제, 기대순위, 1위 수용도)로 돌려준다.

    mcda.evaluate_mcda + rank_smaa2 와 같은 집계다. 상위 K만 필요해도 전체를 계산한다 —
    정답이 몇 위였는지 보려면 끝까지 있어야 한다.
    """
    if arm["vectors"] is None:  # 원본은 축 구성·유틸리티가 달라 그 모듈을 그대로 쓴다
        return _scores_original(candidates, profile)

    criteria = arm["criteria"]
    utilities = utility_rows(candidates, criteria, arm["transform"])
    weights = [boosted(v, criteria, profile.get("priorities"), profile.get("comparison_goals"))
               for v in arm["vectors"]]
    n = len(candidates)
    rank_sum = [0] * n
    first = [0] * n
    for weight in weights:
        totals = [sum(x * u for x, u in zip(weight, row)) for row in utilities]
        order = sorted(range(n), key=lambda i: (-totals[i], str(candidates[i]["plan_id"])))
        first[order[0]] += 1
        for rank, index in enumerate(order):
            rank_sum[index] += rank + 1
    order = sorted(range(n), key=lambda i: (rank_sum[i], str(candidates[i]["plan_id"])))
    return [(candidates[i], rank_sum[i] / len(weights), first[i] / len(weights)) for i in order]


def _scores_original(candidates: list[dict], profile_dict: dict) -> list[tuple]:
    """감마 난수 가중치 SMAA-2(eb0e8a2 이전) 그대로. 축이 10개라 유틸리티도 다르다."""
    profile = UserProfile(**profile_dict)
    reference, _ = _resolve_reference(profile)
    decisions = mcda_original.rank_smaa2(mcda_original.evaluate_mcda(
        candidates, profile.priorities, profile.mvno_brand or profile.host_mno,
        reference=reference, comparison_goals=profile.comparison_goals,
    ))
    by_id = {str(c["plan_id"]): c for c in candidates}
    return [(by_id[d.plan_id], d.smaa2_expected_rank, d.smaa2_first_rank_acceptability)
            for d in decisions]


def rank_top_k(candidates: list[dict], arm: dict, profile: dict) -> list[dict]:
    return [_item(plan) for plan, _, _ in smaa_scores(candidates, arm, profile)[:TOP_K]]


def rank_all(candidates: list[dict], arm: dict, profile: dict) -> list[dict]:
    return [_item(plan) for plan, _, _ in smaa_scores(candidates, arm, profile)]


CONDITION_LABEL = {
    "budget_min_won": "예산 하한", "budget_max_won": "예산 상한", "min_data_gb": "데이터 최소",
    "data_unlimited": "데이터 무제한", "min_qos_mbps": "QoS 최소", "min_tethering_gb": "테더링 최소",
    "min_voice_minutes": "통화 최소", "voice_unlimited": "통화 무제한", "sms_unlimited": "문자 무제한",
    "carrier_type": "통신사 구분", "host_mno": "망", "mvno_brand": "브랜드", "network_gen": "세대",
    "age_condition": "연령", "wanted_benefits": "원하는 혜택", "min_discount_period_months": "할인 기간",
}


def _condition_summary(profile: dict) -> tuple[str, dict]:
    """프로필에서 실제로 걸린 조건만 추려 한 줄 요약과 JSON 을 만든다."""
    used = {k: v for k, v in profile.items()
            if k in CONDITION_LABEL and v not in (None, [], "", False)}
    parts = []
    for key, value in used.items():
        if isinstance(value, bool):
            parts.append(CONDITION_LABEL[key])
        elif isinstance(value, (int, float)) and "won" in key:
            parts.append(f"{CONDITION_LABEL[key]} {value:,}원")
        else:
            parts.append(f"{CONDITION_LABEL[key]} {value}")
    return " · ".join(parts), used


def _plan_line(rank: int, plan: dict, expected: float, acceptability: float) -> str:
    data = "무제한" if plan.get("data_unlimited") else f"{plan.get('data_gb') or 0:g}GB"
    voice = "통화 무제한" if plan.get("voice_unlimited") else f"통화 {plan.get('voice_minutes') or 0:g}분"
    fee = plan.get("discounted_fee") or plan.get("monthly_fee") or 0
    return (f"{rank}. {plan['plan_name']} [기대순위 {expected:.2f} · 1위 {acceptability:.0%}] "
            f"{fee:,}원 · {data} · {voice}")


def write_benchmark(questions, answers, frozen, profiles, arm, arm_key):
    """test2/data/benchmark 와 같은 3종 산출물. 리포트·평가는 이 하네스가 안 태워 비운다."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    OUT_BENCH.mkdir(parents=True, exist_ok=True)
    trace_wb = Workbook()
    ws = trace_wb.active
    ws.title = "결과"
    ws.append(["no", "id", "level", "시나리오", "조건요약", "프로필_JSON", "후보수",
               "추천요금제", "정답요금제", "맞은개수", "정답판정", "에러"])

    pr = Workbook()
    ws_pr = pr.active
    ws_pr.title = "문항별 결과"
    ws_pr.append(["id", "level", "question", "matched", "ans_count", "chat_count",
                  "precision", "recall", "chat_empty"])

    wide = Workbook()
    ws_wide = wide.active
    ws_wide.title = SHEET
    header = ["id", "level", "question"]
    for i in range(1, TOP_K + 1):
        header += [f"ans{i}_name", f"ans{i}_monthly_fee", f"ans{i}_discounted_fee", f"ans{i}_carrier_type"]
    ws_wide.append(header)

    lines = []
    for no, (qid, level, question) in enumerate(questions, start=1):
        candidates = frozen[qid]
        scored = smaa_scores(candidates, arm, profiles[qid]) if candidates else []
        picks = scored[:TOP_K]
        items = [_item(p) for p, _, _ in picks]
        ans_items = answers[qid]["items"]
        matched = count_intersection(ans_items, items)

        verdicts = []
        for ans in ans_items:
            place = next((j for j, (plan, _, _) in enumerate(scored, start=1)
                          if count_intersection([ans], [_item(plan)])), None)
            verdicts.append("맞음" if place and place <= TOP_K
                            else (f"{place}위로 밀림" if place else "후보군에 없음"))
        summary, used = _condition_summary(profiles[qid])

        ws.append([
            no, qid, level, question, summary, json.dumps(used, ensure_ascii=False),
            len(candidates),
            "\n".join(_plan_line(i, p, e, a) for i, (p, e, a) in enumerate(picks, start=1)),
            "\n".join(f"{i}. {a['name']} {a['price']:,}원" for i, a in enumerate(ans_items, start=1)),
            matched,
            "\n".join(f"{i}. {v}" for i, v in enumerate(verdicts, start=1)),
            "" if candidates else "후보 0개",
        ])

        ws_pr.append([qid, level, question, matched, len(ans_items), len(items),
                      round(matched / len(items), 4) if items else 0.0,
                      round(matched / len(ans_items), 4) if ans_items else 0.0,
                      not items])

        row = [qid, level, question]
        for i in range(TOP_K):
            plan = picks[i][0] if i < len(picks) else None
            row += ([plan["plan_name"], plan.get("monthly_fee"), plan.get("discounted_fee"),
                     plan.get("carrier_type")] if plan else [None] * 4)
        ws_wide.append(row)

        lines.append({
            "id": qid, "level": level, "question": question, "arm": arm_key,
            "conditions": used, "candidate_count": len(candidates),
            "picks": [{"rank": i, "plan_id": p["plan_id"], "plan_name": p["plan_name"],
                       "fee": p.get("discounted_fee") or p.get("monthly_fee"),
                       "expected_rank": round(e, 3), "first_rank_acceptability": round(a, 3)}
                      for i, (p, e, a) in enumerate(picks, start=1)],
            "answers": [{"rank": i, "plan_name": a["name"], "price": a["price"], "verdict": v}
                        for i, (a, v) in enumerate(zip(ans_items, verdicts), start=1)],
            "matched": matched,
        })

    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFGHIJKL", [5, 9, 6, 38, 34, 30, 8, 66, 46, 8, 24, 10]):
        ws.column_dimensions[col].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws_sum = pr.create_sheet("레벨별 요약")
    ws_sum.append(["level", "count", "avg_precision", "avg_recall", "empty_count"])
    scored_rows = list(ws_pr.iter_rows(min_row=2, values_only=True))
    for level in sorted({r[1] for r in scored_rows}, key=str):
        group = [r for r in scored_rows if r[1] == level]
        ws_sum.append([level, len(group),
                       round(sum(r[6] for r in group) / len(group), 4),
                       round(sum(r[7] for r in group) / len(group), 4),
                       sum(1 for r in group if r[8])])
    ws_sum.append([])
    ws_sum.append(["전체", len(scored_rows),
                   round(sum(r[6] for r in scored_rows) / len(scored_rows), 4),
                   round(sum(r[7] for r in scored_rows) / len(scored_rows), 4),
                   sum(1 for r in scored_rows if r[8])])

    pr.save(OUT_BENCH / "precision_recall.xlsx")
    trace_wb.save(OUT_BENCH / "agent_trace.xlsx")
    wide.save(OUT_BENCH / "recommend_results.xlsx")
    (OUT_BENCH / "agent_trace.jsonl").write_text(
        "\n".join(json.dumps(line, ensure_ascii=False) for line in lines), encoding="utf-8")
    print(f"저장: {OUT_BENCH}/ (agent_trace.xlsx · agent_trace.jsonl · "
          f"precision_recall.xlsx · recommend_results.xlsx)")


def write_detail(questions, answers, frozen, profiles, arm, arm_key):
    """문항별 정답 5개가 우리 순위에서 몇 위였는지, 우리 top5가 정답이었는지."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "정답 기준"
    ws.append(["id", "level", "question", "정답 순위", "정답 요금제", "정답 가격",
               "우리 순위", "판정"])
    ws2 = wb.create_sheet("추천 기준")
    ws2.append(["id", "level", "question", "우리 순위", "추천 요금제", "추천 가격", "정답 여부"])

    for qid, level, question in questions:
        candidates = frozen[qid]
        ranked = rank_all(candidates, arm, profiles[qid]) if candidates else []
        picks = ranked[:TOP_K]

        for i, ans in enumerate(answers[qid]["items"], start=1):
            place = next((j for j, r in enumerate(ranked, start=1)
                          if count_intersection([ans], [r])), None)
            if place is None:
                verdict = "후보군에 없음 (필터)"
            elif place <= TOP_K:
                verdict = "맞음"
            else:
                verdict = "후보엔 있으나 밀림 (랭킹)"
            ws.append([qid, level, question, i, ans["name"], ans["price"],
                       place if place else "", verdict])

        for i, pick in enumerate(picks, start=1):
            hit = bool(count_intersection(answers[qid]["items"], [pick]))
            ws2.append([qid, level, question, i, pick["name"], pick["price"],
                        "정답" if hit else "오답"])

    for sheet, widths in ((ws, [9, 6, 42, 8, 34, 10, 8, 22]), (ws2, [9, 6, 42, 8, 34, 10, 8])):
        sheet.freeze_panes = "A2"
        for col, width in zip("ABCDEFGH", widths):
            sheet.column_dimensions[col].width = width
        for cell in sheet[1]:
            cell.font = Font(bold=True)

    wb.save(OUT_DETAIL)

    counts = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        counts[row[7]] = counts.get(row[7], 0) + 1
    total = sum(counts.values())
    print(f"\n정답 {total}건 판정 ({arm_key})")
    for name, n in sorted(counts.items(), key=lambda kv: -kv[1]):
        print(f"  {name:<24} {n:>3}건 ({n / total:.0%})")
    print(f"저장: {OUT_DETAIL}")


def _item(plan: dict) -> dict:
    return {"name": plan["plan_name"],
            "price": plan.get("discounted_fee") or plan.get("monthly_fee"),
            "carrier": plan.get("carrier_type") or ""}


def main():
    sys.stdout.reconfigure(encoding="utf-8")  # 윈도우 콘솔이 cp949 라 한글 표가 깨진다
    ap = argparse.ArgumentParser()
    ap.add_argument("--refresh", action="store_true", help="프로파일·가중치 캐시 둘 다 재생성")
    # 가중치만 다시 뽑을 때 프로파일까지 재생성하면 후보군이 바뀌어 arm 비교가 깨진다
    ap.add_argument("--refresh-arms", action="store_true", help="가중치 캐시만 재생성")
    ap.add_argument("--workers", type=int, default=4)
    ap.add_argument("--detail", default="D_ridge6_tf",
                    help="문항별 정답 대조표를 뽑을 arm (빈 문자열이면 생략)")
    args = ap.parse_args()

    questions = load_questions()
    answers = load_wide_rows(str(ANSWER_KEY))
    profiles = load_profiles(questions, args.refresh, args.workers)
    arms = load_arms(args.refresh or args.refresh_arms)

    frozen = {qid: candidates_of(profiles[qid]) for qid, _, _ in questions}
    print(f"\n후보군 동결 · 후보 0개 문항 {sum(1 for v in frozen.values() if not v)}개")

    rows = []
    # 상한: 정답이 애초에 후보군에 들어 있는가. 가중치로는 못 넘는 천장이다.
    for qid, level, _ in questions:
        items = [_item(c) for c in frozen[qid]]
        hit = count_intersection(answers[qid]["items"], items)
        # 후보가 K개 미만이면 arm 도 그만큼만 내놓는다. 분모를 맞춰야 비교가 된다.
        slots = min(TOP_K, len(items))
        rows.append({"id": qid, "level": level, "arm": "ceiling",
                     "precision": min(hit, slots) / slots if slots else 0.0,
                     "recall": hit / len(answers[qid]["items"]) if answers[qid]["items"] else 0.0,
                     "empty": not items})

    for key, arm in arms.items():
        for qid, level, _ in questions:
            picks = rank_top_k(frozen[qid], arm, profiles[qid]) if frozen[qid] else []
            matched = count_intersection(answers[qid]["items"], picks)
            rows.append({
                "id": qid, "level": level, "arm": key,
                "precision": matched / len(picks) if picks else 0.0,
                "recall": matched / len(answers[qid]["items"]) if answers[qid]["items"] else 0.0,
                "empty": not picks,
            })

    df = pd.DataFrame(rows)
    df.to_csv(OUT_CSV, index=False, encoding="utf-8-sig")

    labels = {k: v["label"] for k, v in arms.items()} | {"ceiling": "후보군 상한"}
    overall = df.groupby("arm")[["precision", "recall"]].mean().round(3)
    overall["빈 응답"] = df.groupby("arm")["empty"].sum()
    overall["설명"] = pd.Series(labels)
    print("\n전체 (macro-avg, 100문항)")
    print(overall.to_string())

    print("\n레벨별")
    print(df.pivot_table(index="arm", columns="level", values=["precision", "recall"]).round(3).to_string())

    print("\n빈 응답 제외")
    live = df[~df["empty"]]
    print(live.groupby("arm")[["precision", "recall"]].mean().round(3).to_string())
    print(f"\n저장: {OUT_CSV}")

    if args.detail:
        write_detail(questions, answers, frozen, profiles, arms[args.detail], args.detail)
        write_benchmark(questions, answers, frozen, profiles, arms[args.detail], args.detail)


if __name__ == "__main__":
    main()
